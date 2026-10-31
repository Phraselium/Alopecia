"""Genera ranking.xlsx (hojas Anuncios, Top20, Supuestos) con el scoring en fórmulas de Excel.

Convención de colores:
  azul  (0000FF) = dato introducido a mano o scrapeado (con su fuente en nota o columna)
  negro (000000) = calculado con fórmulas de la propia hoja
  verde (008000) = fórmula que trae datos de otra hoja (incluye nombres definidos de "Supuestos")
No se inmovilizan paneles en ninguna hoja.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

from .models import Listing
from .scoring import Params

BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
HDR_FILL = PatternFill("solid", fgColor="DDE4EE")
SEC_FONT = Font(bold=True, size=12)


def _font(color, bold=False):
    return Font(color=color, bold=bold)


def _yn(v):
    return None if v is None else ("sí" if v else "no")


class Names:
    """Registro de nombres definidos -> se usan en fórmulas de otras hojas."""

    def __init__(self, wb: Workbook):
        self.wb = wb
        self.names: set[str] = set()

    def add(self, name: str, ref: str):
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
        self.names.add(name)


# =====================================================================================
# Supuestos
# =====================================================================================
def build_supuestos(wb: Workbook, p: Params, names: Names) -> None:
    ws = wb.create_sheet("Supuestos")
    cfg, fis, rel = p.cfg, p.fiscal, p.rel
    sc = cfg["scoring"]
    r = 1
    ws.cell(r, 1, "Supuestos: pesos, parámetros, fiscalidad y fiabilidad (cada dato con su estado y fuente)").font = SEC_FONT
    r += 1
    ws.cell(r, 1, "Azul = dato introducido/fuente externa · Negro = calculado en esta hoja · Verde = viene de otra hoja. "
                  "Estado: confirmado / probable (verificar en la fuente) / PENDIENTE (no se usa como cierto).")
    r += 2

    def header(cols):
        nonlocal r
        for i, h in enumerate(cols, 1):
            c = ws.cell(r, i, h)
            c.font = Font(bold=True)
            c.fill = HDR_FILL
        r += 1

    def param(name, label, value, estado="criterio (config.yaml)", fuente="config.yaml", formula=False):
        nonlocal r
        ws.cell(r, 1, label)
        c = ws.cell(r, 2, value)
        c.font = _font(BLACK if formula else BLUE)
        ws.cell(r, 3, estado)
        ws.cell(r, 4, fuente)
        names.add(name, f"Supuestos!$B${r}")
        r += 1

    def section(t):
        nonlocal r
        r += 1
        ws.cell(r, 1, t).font = SEC_FONT
        r += 1

    # ---- pesos
    section("1. Pesos del scoring (deben sumar 1)")
    header(["Criterio", "Peso", "Estado", "Fuente"])
    w = sc["pesos"]
    first = r
    param("W_FIAB", "Fiabilidad modelo/motor", w["fiabilidad"])
    param("W_PRECIO", "Precio frente al mercado", w["precio"])
    param("W_KMEDAD", "Km y edad", w["km_edad"])
    param("W_ANUNCIO", "Calidad del anuncio / mantenimiento", w["anuncio"])
    param("W_USO", "Coste de uso en Valencia", w["coste_uso"])
    param("W_SEGURIDAD", "Seguridad de compra (garantía, km certificados, revisión, devolución, historial)", w["seguridad"],
          "criterio (config.yaml)", "config.yaml: añadido el 31-10-2026; los 5 pesos anteriores × 0,85 redondeados")
    param("W_SUMA", "Suma de pesos (control)", f"=SUM(B{first}:B{r-1})", "calculado", "", formula=True)

    # ---- parámetros
    section("2. Parámetros de búsqueda y de puntuación")
    header(["Parámetro", "Valor", "Estado", "Fuente"])
    f, k, a, u = cfg["filtros"], sc["km_edad"], sc["anuncio"], sc["coste_uso"]
    param("ANIO_ACTUAL", "Año de referencia", fis["anio_actual"], "dato", "fiscal.yaml")
    param("ANIO_MIN", "Año mínimo", f["anio_min"])
    param("KM_MAX", "Km máximos", f["km_max"])
    param("PRECIO_MAX", "Precio anuncio máximo (€)", cfg["presupuesto"]["precio_anuncio_max_eur"])
    param("TOTAL_MAX", "Coste total máximo (€)", cfg["presupuesto"]["total_max_eur"])
    param("SEG_ANIO", "Tramo de años para mediana", cfg["segmentos"]["anio_tramo"])
    param("SEG_KM", "Tramo de km para mediana", cfg["segmentos"]["km_tramo"])
    param("MIN_N", "Mín. anuncios por segmento (si no, se sube de nivel)", cfg["segmentos"]["min_muestras"])
    param("PRECIO_PEND", "Puntos por unidad de descuento (50 + desc·k)", sc["precio"]["pendiente"])
    param("KMA_IDEAL_MAX", "Km/año ideal máximo (100 puntos hasta aquí)", k["km_anio_ideal_max"])
    param("KMA_CERO", "Km/año con 0 puntos", k["km_anio_cero"])
    param("KMA_SOSP", "Km/año sospechosamente bajo", k["km_anio_sospechoso"])
    param("KMA_PEN", "Puntos km/año si es sospechoso y no justificado", k["penaliza_sospechoso"])
    param("W_KMT", "Peso km totales dentro de Km y edad", k["peso_km_total"])
    param("W_KMA", "Peso km/año dentro de Km y edad", k["peso_km_anio"])
    param("W_EDAD", "Peso edad dentro de Km y edad", k["peso_edad"])
    param("P_ITV", "Puntos: ITV en vigor", a["itv"])
    param("P_LIBRO", "Puntos: libro de mantenimiento", a["libro"])
    param("P_CORREA", "Puntos: correa/distribución cambiada", a["correa"])
    param("P_PROP", "Puntos: un solo propietario", a["propietario_unico"])
    param("P_GAR", "Puntos: garantía", a["garantia"])
    param("P_FOTOS", "Puntos: fotos (proporcional)", a["fotos"])
    param("FOTOS_OBJ", "Fotos para puntuación completa", a["fotos_objetivo"])
    param("W_ET", "Peso etiqueta dentro de Coste de uso", u["peso_etiqueta"])
    param("W_CONS", "Peso consumo dentro de Coste de uso", u["peso_consumo"])
    param("W_SEG", "Peso seguro dentro de Coste de uso", u["peso_seguro"])
    param("NEUTRO", "Puntuación neutra cuando el dato está PENDIENTE", u["neutro_si_pendiente"])
    sg = sc["seguridad"]
    gl = fis["garantia_legal"]
    param("SG_GAR", "Seguridad: puntos por garantía (proporcional a meses)", sg["garantia"])
    param("SG_GAR_FULL", "Seguridad: meses de garantía para puntuación completa", sg["garantia_meses_full"])
    param("SG_AMPL", "Seguridad: puntos si la garantía es ampliable", sg["ampliable"])
    param("SG_KMC", "Seguridad: puntos por km certificados", sg["km_certificados"])
    param("SG_REV", "Seguridad: puntos por revisión documentada", sg["revision"])
    param("SG_DEV", "Seguridad: puntos por derecho de devolución (proporcional a días)", sg["devolucion"])
    param("SG_DEV_FULL", "Seguridad: días de devolución para puntuación completa", sg["devolucion_dias_full"])
    param("SG_HIST", "Seguridad: puntos por historial/libro de mantenimiento", sg["historial"])
    param("SG_LEGAL_ON", "Aplicar mínimo legal a profesionales sin garantía declarada (1 = sí)",
          1 if sg.get("aplicar_minimo_legal_profesional", True) else 0)
    param("GAR_LEGAL", "Garantía legal mínima en usados de vendedor profesional (meses)", gl["minimo_meses_segunda_mano"],
          gl["estado"], f"{gl['norma']} — {gl['fuente']}")

    # ---- fiscalidad
    section("3. Coste de compra: ITP (Comunitat Valenciana) y tasas DGT")
    header(["Concepto", "Valor", "Estado", "Fuente"])
    itp = fis["itp"]
    param("ITP_TIPO", "ITP tipo general (si no aplica cuota fija)", itp["tipo"], itp["tipo_estado"], itp["tipo_fuente"])
    param("CUOTA_VALOR_MAX", "Cuota fija solo si valor < (€)", itp["cuota_fija_valor_max_eur"], itp["cuotas_fijas_estado"], itp["cuotas_fijas_fuente"])
    param("TASA_DGT", "Tasa DGT transferencia 1.5 (€)", fis["dgt"]["tasa_transferencia_eur"], fis["dgt"]["tasa_transferencia_estado"], fis["dgt"]["tasa_transferencia_fuente"])
    param("INFORME_DGT", "Informe DGT 4.1 (€)", fis["dgt"]["informe_vehiculo_eur"], fis["dgt"]["informe_vehiculo_estado"], fis["dgt"]["informe_vehiculo_fuente"])
    tramos = sorted({c["antiguedad_mas_de_anios"] for c in itp["cuotas_fijas"]}, reverse=True)
    param("ITP_TRAMO_ALTO", "Cuota fija: tramo 'más de N años' (alto)", tramos[0], itp["cuotas_fijas_estado"], itp["cuotas_fijas_fuente"])
    param("ITP_TRAMO_BAJO", "Cuota fija: tramo 'más de N años' (bajo)", tramos[-1], itp["cuotas_fijas_estado"], itp["cuotas_fijas_fuente"])
    ws.cell(r, 1, "Vendedor profesional: sin ITP (IVA/REBU incluido); se suman sus gastos de gestión y la tasa DGT solo si "
                  "esos gastos no incluyen la transferencia o no constan — " + itp["exento_si_profesional_estado"])
    r += 1
    ws.cell(r, 1, "Nota: " + itp["nota_ley_5_2025"])
    r += 2
    ws.cell(r, 1, "Tabla de cuotas fijas de ITP").font = Font(bold=True)
    r += 1
    header(["Más de (años)", "cc desde", "cc hasta", "Cuota (€)", "Descripción", "Estado", "Fuente"])
    t0 = r
    for c in itp["cuotas_fijas"]:
        for i, v in enumerate([c["antiguedad_mas_de_anios"], c["cc_min"], c["cc_max"] or 99999, c["cuota_eur"]], 1):
            ws.cell(r, i, v).font = _font(BLUE)
        ws.cell(r, 5, c["descripcion"])
        ws.cell(r, 6, itp["cuotas_fijas_estado"])
        ws.cell(r, 7, itp["cuotas_fijas_fuente"])
        r += 1
    for col, nm in zip("ABCD", ["CF_MAS", "CF_CCMIN", "CF_CCMAX", "CF_CUOTA"]):
        names.add(nm, f"Supuestos!${col}${t0}:${col}${r-1}")
    ws.cell(r, 1, "cc hasta 99999 = sin límite superior. Cilindrada desconocida -> se usa 99999 (tramo más caro).")
    r += 1
    pm = fis["precios_medios"]
    ws.cell(r, 1, f"Valores fiscales de Hacienda: {pm['norma']} ({pm['fuente']}). Estado: {pm['estado']}. "
                  f"Solo afectan a coches de 5 años o menos (tipo porcentual).")
    r += 1

    # ---- etiqueta y ZBE
    section("4. Etiqueta DGT y ZBE de València")
    header(["Concepto", "Valor", "Estado", "Fuente"])
    e = fis["etiqueta_dgt"]
    param("GAS_C", "Gasolina etiqueta C desde (año)", e["gasolina_C_desde"], e["estado"], e["fuente"])
    param("GAS_B", "Gasolina etiqueta B desde (año)", e["gasolina_B_desde"], e["estado"], e["fuente"])
    param("DIE_C", "Diésel etiqueta C desde (año)", e["diesel_C_desde"], e["estado"], e["fuente"])
    param("DIE_B", "Diésel etiqueta B desde (año)", e["diesel_B_desde"], e["estado"], e["fuente"])
    z = fis["zbe_valencia"]
    ws.cell(r, 1, "ZBE València")
    ws.cell(r, 2, z["resumen"].strip()).alignment = Alignment(wrap_text=True)
    ws.cell(r, 3, z["estado"])
    ws.cell(r, 4, z["fuente"])
    r += 2
    header(["Etiqueta", "Puntos (criterio)", "Estado", "Fuente"])
    t0 = r
    for et, pts in u["puntos_etiqueta"].items():
        ws.cell(r, 1, et).font = _font(BLUE)
        ws.cell(r, 2, pts).font = _font(BLUE)
        ws.cell(r, 3, "criterio (config.yaml)")
        ws.cell(r, 4, "config.yaml; la ZBE solo restringe sin etiqueta (A), ver fila ZBE")
        r += 1
    names.add("ETQ_TBL", f"Supuestos!$A${t0}:$B${r-1}")
    ws.cell(r, 1, "Consumo (l/100 km): PENDIENTE de fuente -> se puntúa con NEUTRO.")
    r += 2
    se = fis["seguro_estimado"]
    ws.cell(r, 1, "Seguro estimado (€/año) — " + se["descripcion"]).font = Font(bold=True)
    r += 1
    header(["Segmento", "Mínimo", "Estimación central", "Máximo", "Estado", "Fuentes"])
    t0 = r
    for segm, v in se["segmentos"].items():
        ws.cell(r, 1, segm).font = _font(BLUE)
        for i, k in enumerate(["min", "central", "max"], 2):
            ws.cell(r, i, v[k]).font = _font(BLUE)
        ws.cell(r, 5, se["estado"])
        ws.cell(r, 6, " | ".join(se["fuentes"]))
        r += 1
    names.add("SEGURO_TBL", f"Supuestos!$A${t0}:$D${r-1}")
    ws.cell(r, 1, "Nota: " + " ".join(se["notas"].split()))
    r += 1
    param("SEG_REF_MIN", "Seguro: referencia para 100 puntos (mínimo de la tabla)", f"=MIN(B{t0}:B{t0 + len(se['segmentos']) - 1})",
          "calculado", "", formula=True)
    param("SEG_REF_MAX", "Seguro: referencia para 0 puntos (máximo de la tabla)", f"=MAX(D{t0}:D{t0 + len(se['segmentos']) - 1})",
          "calculado", "", formula=True)

    # ---- fiabilidad
    section("5. Fiabilidad: medias TÜV por franja de edad")
    header(["Franja", "Mängelquote media (%)", "Estado", "Fuente"])
    t0 = r
    for fr, d in rel["medias_tuv_2026"].items():
        ws.cell(r, 1, fr).font = _font(BLUE)
        ws.cell(r, 2, d["valor"]).font = _font(BLUE)
        ws.cell(r, 3, d["estado"])
        ws.cell(r, 4, d["fuente"])
        r += 1
    names.add("TUV_MEDIAS", f"Supuestos!$A${t0}:$B${r-1}")

    section("6. Fiabilidad: índice OCU 2026 por marca")
    ocu = rel["marcas_ocu_2026"]
    header(["Marca", "Índice OCU (0-100)", "Estado", "Fuente"])
    t0 = r
    for m, v in ocu["valores"].items():
        ws.cell(r, 1, m).font = _font(BLUE)
        ws.cell(r, 2, v).font = _font(BLUE)
        ws.cell(r, 3, ocu["estado"])
        ws.cell(r, 4, ocu["fuente"])
        r += 1
    names.add("OCU_TBL", f"Supuestos!$A${t0}:$B${r-1}")

    section("6b. Fiabilidad: ADAC Pannenstatistik 2025 — averías por 1.000 vehículos, matriculación 2015")
    ad = rel["adac_2025"]
    header(["Fila ADAC", "Pannenkennziffer EZ 2015", "Estado", "Fuente"])
    t0 = r
    for k, v in ad["pkz_2015"].items():
        ws.cell(r, 1, k).font = _font(BLUE)
        ws.cell(r, 2, v).font = _font(BLUE)
        ws.cell(r, 3, ad["estado"] if r == t0 else "")
        ws.cell(r, 4, ad["fuente"] if r == t0 else "")
        r += 1
    names.add("ADAC_TBL", f"Supuestos!$A${t0}:$B${r-1}")
    param("ADAC_REF", "ADAC: mediana de la tabla (referencia 50 puntos)", f"=MEDIAN(B{t0}:B{r-1})", "calculado", ad["puntuacion"], formula=True)

    section("7. Fiabilidad por modelo (punt. TÜV = 50 + 100·(media − MQ)/media; punt. ADAC = 50 + 100·(ref − PKZ)/ref; punt. modelo = media de TÜV, ADAC y OCU marca)")
    header(["Modelo", "Marca", "TÜV Mängelquote (%)", "Franja", "Media franja (%)", "Punt. TÜV",
            "Fila ADAC", "ADAC PKZ 2015", "Punt. ADAC", "Punt. OCU marca", "Punt. modelo", "Pendiente", "Estado", "Fuente"])
    t0 = r
    for key, m in rel["modelos"].items():
        marca = key.split(" ")[0]
        ws.cell(r, 1, key).font = _font(BLUE)
        ws.cell(r, 2, marca).font = _font(BLUE)
        if m.get("tuv_mq") is not None:
            ws.cell(r, 3, m["tuv_mq"]).font = _font(BLUE)
            ws.cell(r, 4, m["tuv_franja"]).font = _font(BLUE)
            if m.get("tuv_media_propia"):
                ws.cell(r, 5, m["tuv_media_propia"]).font = _font(BLUE)
            else:
                ws.cell(r, 5, f"=VLOOKUP(D{r},TUV_MEDIAS,2,0)").font = _font(BLACK)
        ws.cell(r, 6, f'=IF(C{r}="","",MAX(0,MIN(100,50+100*(E{r}-C{r})/E{r})))').font = _font(BLACK)
        if m.get("adac"):
            ws.cell(r, 7, m["adac"]).font = _font(BLUE)
        ws.cell(r, 8, f'=IF(G{r}="","",IFERROR(VLOOKUP(G{r},ADAC_TBL,2,0),""))').font = _font(BLACK)
        ws.cell(r, 9, f'=IF(H{r}="","",MAX(0,MIN(100,50+100*(ADAC_REF-H{r})/ADAC_REF)))').font = _font(BLACK)
        ws.cell(r, 10, f'=IFERROR(VLOOKUP(B{r},OCU_TBL,2,0),"")').font = _font(BLACK)
        ws.cell(r, 11, f'=IF(COUNT(F{r},I{r},J{r})=0,NEUTRO,AVERAGE(F{r},I{r},J{r}))').font = _font(BLACK)
        ws.cell(r, 12, f'=IF(COUNT(F{r},I{r},J{r})=0,"PENDIENTE","")').font = _font(BLACK)
        ws.cell(r, 13, m.get("estado"))
        ws.cell(r, 14, m.get("fuente", ""))
        r += 1
    names.add("MOD_TBL", f"Supuestos!$A${t0}:$N${r-1}")

    section("8. Ajustes por motor con problemas conocidos (el problema tiene fuente; la magnitud es criterio)")
    header(["Marca", "Motor", "Año desde", "Año hasta", "Ajuste (puntos)", "Problema", "Qué revisar", "Estado", "Fuente"])
    t0 = r
    for m in rel["motores"]:
        for i, key in enumerate(["marca", "motor", "anio_desde", "anio_hasta", "ajuste"], 1):
            ws.cell(r, i, m[key]).font = _font(BLUE)
        ws.cell(r, 6, m["problema"])
        ws.cell(r, 7, m["revisar"])
        ws.cell(r, 8, m["estado"])
        ws.cell(r, 9, m["fuente"])
        r += 1
    for col, nm in zip("ABCDE", ["MOT_MARCA", "MOT_MOTOR", "MOT_DESDE", "MOT_HASTA", "MOT_AJUSTE"]):
        names.add(nm, f"Supuestos!${col}${t0}:${col}${r-1}")
    ws.cell(r, 1, "Motores con buena fama: PENDIENTE de fuente de autoridad; sin ajuste positivo.")
    r += 1

    section("9. Ajustes por combustible")
    header(["Combustible", "Ajuste (puntos)", "Problema", "Qué revisar", "Estado", "Fuente"])
    t0 = r
    for c in rel["combustibles"]:
        ws.cell(r, 1, c["combustible"]).font = _font(BLUE)
        ws.cell(r, 2, c["ajuste"]).font = _font(BLUE)
        ws.cell(r, 3, c["problema"])
        ws.cell(r, 4, c["revisar"])
        ws.cell(r, 5, c["estado"])
        ws.cell(r, 6, c["fuente"])
        r += 1
    names.add("COMB_COMB", f"Supuestos!$A${t0}:$A${r-1}")
    names.add("COMB_AJUSTE", f"Supuestos!$B${t0}:$B${r-1}")

    for col, wdt in zip("ABCDEFGHIJK", [44, 18, 22, 22, 18, 40, 40, 14, 12, 30, 60]):
        ws.column_dimensions[col].width = wdt


# =====================================================================================
# Anuncios
# =====================================================================================
# (clave, cabecera, tipo, plantilla). tipo: 'data' (azul), 'calc' (negro/verde según refs), 'array'
COLUMNS = [
    ("uid", "ID", "data", None), ("plataforma", "Plataforma", "data", None), ("url", "URL", "data", None),
    ("titulo", "Título", "data", None), ("marca", "Marca", "data", None), ("modelo_key", "Modelo", "data", None),
    ("motor", "Motor (normalizado)", "data", None), ("version", "Versión", "data", None),
    ("cc", "Cilindrada (cc)", "data", None), ("cv", "Potencia (CV)", "data", None),
    ("combustible", "Combustible", "data", None), ("carroceria", "Segmento", "data", None), ("anio", "Año", "data", None), ("km", "Km", "data", None),
    ("precio", "Precio (€)", "data", None), ("cambio", "Cambio", "data", None),
    ("automatico", "Automático", "data", None), ("vendedor", "Vendedor", "data", None),
    ("ubicacion", "Ubicación", "data", None), ("fecha", "Publicado", "data", None),
    ("fotos", "Nº fotos", "data", None), ("itv", "ITV en vigor", "data", None),
    ("libro", "Libro mantenimiento", "data", None), ("correa", "Correa/distrib. cambiada", "data", None),
    ("propietarios", "Nº propietarios", "data", None), ("garantia", "Garantía", "data", None),
    ("km_just", "Km justificados", "data", None),
    ("gar_meses", "Garantía incluida (meses)", "data", None), ("gar_ampl", "Garantía ampliable", "data", None),
    ("gar_ampl_eur", "Precio ampliación garantía (€)", "data", None), ("revision", "Revisión (puntos)", "data", None),
    ("km_cert", "Km certificados", "data", None), ("devolucion", "Devolución (días)", "data", None),
    ("p_contado", "Precio al contado (€)", "data", None), ("p_financiado", "Precio financiado (€)", "data", None),
    ("solo_fin", "Precio anunciado solo financiando", "data", None),
    ("gastos", "Gastos de gestión del vendedor (€)", "data", None), ("gastos_transf", "Gastos incluyen transferencia", "data", None),
    ("distancia", "Distancia a València (km)", "data", None),
    ("med_merc", "Mediana de mercado del segmento (€)", "data", None), ("n_merc", "N mercado", "data", None),
    ("nivel_merc", "Nivel mediana mercado", "data", None),
    ("duplicados", "Duplicados fusionados", "data", None),
    ("avisos", "Avisos", "data", None), ("descripcion", "Descripción", "data", None),
    ("fuente", "Fuente (fichero /raw)", "data", None),
    # ---- calculadas
    ("edad", "Edad (años)", "calc", "=MAX(ANIO_ACTUAL-{anio},0)"),
    ("km_anio", "Km/año", "calc", "={km}/MAX({edad},0.5)"),
    ("k1", "Seg. modelo", "calc", "={modelo_key}"),
    ("k2", "Seg. modelo+año", "calc", '={k1}&"|"&INT({anio}/SEG_ANIO)*SEG_ANIO'),
    ("k3", "Seg. modelo+año+km", "calc", '={k2}&"|"&INT({km}/SEG_KM)*SEG_KM'),
    ("n1", "N modelo", "calc", "=COUNTIF({R:k1},{k1})"),
    ("n2", "N modelo+año", "calc", "=COUNTIF({R:k2},{k2})"),
    ("n3", "N modelo+año+km", "calc", "=COUNTIF({R:k3},{k3})"),
    ("mediana", "Mediana segmento (€)", "array",
     "=IF(N({n_merc})>=MIN_N,{med_merc},IF({n3}>=MIN_N,MEDIAN(IF({R:k3}={k3},{R:precio})),IF({n2}>=MIN_N,MEDIAN(IF({R:k2}={k2},{R:precio})),MEDIAN(IF({R:k1}={k1},{R:precio})))))"),
    ("nivel", "Nivel mediana", "calc",
     '=IF(N({n_merc})>=MIN_N,"mercado "&{nivel_merc},IF({n3}>=MIN_N,"modelo+año+km",IF({n2}>=MIN_N,"modelo+año","modelo")))'),
    ("descuento", "Descuento vs mediana", "calc", "=({mediana}-{precio})/{mediana}"),
    ("etiqueta", "Etiqueta DGT", "calc",
     '=IF(OR({combustible}="hibrido",{combustible}="glp"),"ECO",IF({combustible}="electrico","CERO",'
     'IF({combustible}="gasolina",IF({anio}>=GAS_C,"C",IF({anio}>=GAS_B,"B","A")),'
     'IF({combustible}="diesel",IF({anio}>=DIE_C,"C",IF({anio}>=DIE_B,"B","A")),""))))'),
    # fiabilidad
    ("f_base", "Fiab. base modelo", "calc",
     "=IFERROR(VLOOKUP({modelo_key},MOD_TBL,11,0),IFERROR(VLOOKUP({marca},OCU_TBL,2,0),NEUTRO))"),
    ("f_motor", "Ajuste motor", "calc",
     '=SUMIFS(MOT_AJUSTE,MOT_MOTOR,{motor},MOT_MARCA,{marca},MOT_DESDE,"<="&{anio},MOT_HASTA,">="&{anio})'
     '+SUMIFS(MOT_AJUSTE,MOT_MOTOR,{motor},MOT_MARCA,"(todas)",MOT_DESDE,"<="&{anio},MOT_HASTA,">="&{anio})'),
    ("f_comb", "Ajuste combustible", "calc", "=SUMIFS(COMB_AJUSTE,COMB_COMB,{combustible})"),
    ("s_fiab", "P. Fiabilidad", "calc", "=MAX(0,MIN(100,{f_base}+{f_motor}+{f_comb}))"),
    # precio
    ("s_precio", "P. Precio", "calc", "=MAX(0,MIN(100,50+{descuento}*PRECIO_PEND))"),
    # km y edad
    ("s_kmt", "P. km totales", "calc", "=MAX(0,100*(1-{km}/KM_MAX))"),
    ("s_kma", "P. km/año", "calc",
     '=IF({km_anio}<KMA_SOSP,IF({km_just}="sí",100,KMA_PEN),IF({km_anio}<=KMA_IDEAL_MAX,100,'
     "MAX(0,100*(KMA_CERO-{km_anio})/(KMA_CERO-KMA_IDEAL_MAX))))"),
    ("s_edad", "P. edad", "calc", "=MAX(0,MIN(100,100*({anio}-ANIO_MIN)/(ANIO_ACTUAL-ANIO_MIN)))"),
    ("s_kmedad", "P. Km y edad", "calc", "=W_KMT*{s_kmt}+W_KMA*{s_kma}+W_EDAD*{s_edad}"),
    # anuncio
    ("s_anuncio", "P. Anuncio", "calc",
     '=P_ITV*({itv}="sí")+P_LIBRO*({libro}="sí")+P_CORREA*({correa}="sí")+P_PROP*({propietarios}=1)'
     '+P_GAR*({garantia}="sí")+P_FOTOS*MIN(1,N({fotos})/FOTOS_OBJ)'),
    # uso
    ("s_etq", "P. etiqueta", "calc", "=IFERROR(VLOOKUP({etiqueta},ETQ_TBL,2,0),NEUTRO)"),
    ("s_cons", "P. consumo (PENDIENTE)", "calc", "=NEUTRO"),
    ("seguro", "Seguro estimado €/año (aprox.)", "calc", '=IFERROR(VLOOKUP({carroceria},SEGURO_TBL,3,0),"")'),
    ("s_seg", "P. seguro", "calc", '=IF({seguro}="",NEUTRO,MAX(0,MIN(100,100*(SEG_REF_MAX-{seguro})/(SEG_REF_MAX-SEG_REF_MIN))))'),
    ("s_uso", "P. Coste uso", "calc", "=W_ET*{s_etq}+W_CONS*{s_cons}+W_SEG*{s_seg}"),
    # seguridad de compra
    ("gar_ef", "Garantía efectiva (meses)", "calc",
     '=IF(N({gar_meses})>0,{gar_meses},IF(AND({vendedor}="profesional",SG_LEGAL_ON=1),GAR_LEGAL,0))'),
    ("s_seguridad", "P. Seguridad compra", "calc",
     '=SG_GAR*MIN(1,{gar_ef}/SG_GAR_FULL)+SG_AMPL*({gar_ampl}="sí")+SG_KMC*({km_cert}="sí")+SG_REV*(N({revision})>0)'
     '+SG_DEV*MIN(1,N({devolucion})/SG_DEV_FULL)+SG_HIST*({libro}="sí")'),
    ("total", "PUNTUACIÓN", "calc",
     "=W_FIAB*{s_fiab}+W_PRECIO*{s_precio}+W_KMEDAD*{s_kmedad}+W_ANUNCIO*{s_anuncio}+W_USO*{s_uso}+W_SEGURIDAD*{s_seguridad}"),
    # coste
    ("cc_fiscal", "cc para ITP", "calc",
     '=IF(N({cc})>0,{cc},IFERROR(ROUND(_xlfn.NUMBERVALUE(LEFT({motor},FIND(" ",{motor}&" ")-1),".")*1000,0),99999))'),
    ("itp", "ITP (€)", "calc",
     '=IF({vendedor}="profesional",0,IF({precio}<CUOTA_VALOR_MAX,'
     'IF({edad}>ITP_TRAMO_ALTO,SUMIFS(CF_CUOTA,CF_MAS,ITP_TRAMO_ALTO,CF_CCMIN,"<="&{cc_fiscal},CF_CCMAX,">="&{cc_fiscal}),'
     'IF({edad}>ITP_TRAMO_BAJO,SUMIFS(CF_CUOTA,CF_MAS,ITP_TRAMO_BAJO,CF_CCMIN,"<="&{cc_fiscal},CF_CCMAX,">="&{cc_fiscal}),'
     '{precio}*ITP_TIPO)),{precio}*ITP_TIPO))'),
    ("gastos_calc", "Gastos vendedor en coste (€)", "calc", '=IF({vendedor}="profesional",N({gastos}),0)'),
    ("tasa", "Tasa DGT (€)", "calc",
     '=IF(AND({vendedor}="profesional",ISNUMBER({gastos}),{gastos_transf}="sí"),0,TASA_DGT)'),
    ("informe", "Informe DGT (€)", "calc", "=INFORME_DGT"),
    ("coste_total", "COSTE TOTAL (€)", "calc", "={precio}+{gastos_calc}+{itp}+{tasa}+{informe}"),
    ("primer_anio", "Coste 1er año con seguro (€)", "calc", "={coste_total}+N({seguro})"),
    ("apto", "Dentro de presupuesto", "calc", '=IF({coste_total}<=TOTAL_MAX,"sí","no")'),
    ("rank", "Puesto", "calc",
     '=IF({apto}="sí",COUNTIFS({R:apto},"sí",{R:total},">"&{total})+COUNTIFS({RUP:apto},"sí",{RUP:total},{total}),"")'),
]
COL = {key: get_column_letter(i) for i, (key, *_r) in enumerate(COLUMNS, 1)}

NOTES = {
    "uid": "Datos de anuncio (azul): scrapeados de la plataforma; fuente exacta en la columna 'Fuente (fichero /raw)'.",
    "marca": "Normalizada por src/normalize.py a partir del dato scrapeado.",
    "motor": "Normalizado por src/normalize.py (cilindrada + tecnología).",
    "itv": "Extraído del texto del anuncio por src/extract.py: sí / no / vacío = no lo menciona.",
    "mediana": "Si hay MIN_N o más anuncios de mercado del segmento (todos los vistos en los listados, también > precio máx.), "
               "usa la mediana de mercado; si no, fórmula matricial con la mediana de esta hoja (modelo+año+km, subiendo de nivel).",
    "med_merc": "Calculada en Python (src/scoring.py: market_medians) con los anuncios vistos en los listados de las plataformas "
                "(tabla 'mercado' de data/listings.sqlite), mismo esquema de segmentos.",
    "itp": "Cuota fija CV por antigüedad y cilindrada si valor < CUOTA_VALOR_MAX; si no, tipo porcentual. Profesional: 0.",
    "rank": "Desempate por orden de fila.",
}


def _expand(tpl: str, r: int, n_last: int, names: set[str]) -> tuple[str, bool]:
    def rep(m):
        tok = m.group(1)
        if tok.startswith("R:"):
            c = COL[tok[2:]]
            return f"${c}$2:${c}${n_last}"
        if tok.startswith("RUP:"):
            c = COL[tok[4:]]
            return f"${c}$2:{c}{r}"
        return f"{COL[tok]}{r}"
    f = re.sub(r"\{([A-Za-z_:0-9]+)\}", rep, tpl)
    external = any(re.search(rf"(?<![A-Za-z_]){n}(?![A-Za-z_0-9])", f) for n in names)
    return f, external


def listing_row(l: Listing, avisos: str, mkt=None) -> dict:
    mkt = mkt or (None, None, None)
    return {
        "med_merc": mkt[0], "n_merc": mkt[2] if mkt[0] is not None else None, "nivel_merc": mkt[1] or None,
        "uid": l.uid, "plataforma": l.plataforma, "url": l.url, "titulo": l.titulo, "marca": l.marca,
        "modelo_key": l.modelo_key, "motor": l.motor_key, "version": l.version, "cc": l.cilindrada_cc,
        "cv": l.potencia_cv, "combustible": l.combustible, "carroceria": l.carroceria, "anio": l.anio, "km": l.km, "precio": l.precio,
        "cambio": l.cambio, "automatico": "sí" if l.cambio == "automatico" else ("no" if l.cambio else None),
        "vendedor": l.vendedor, "ubicacion": l.ubicacion, "fecha": l.fecha_publicacion, "fotos": l.num_fotos,
        "itv": _yn(l.itv_vigor), "libro": _yn(l.libro_mantenimiento), "correa": _yn(l.correa_cambiada),
        "propietarios": l.num_propietarios, "garantia": _yn(l.garantia), "km_just": _yn(l.km_justificados),
        "gar_meses": l.garantia_meses, "gar_ampl": _yn(l.garantia_ampliable), "gar_ampl_eur": l.garantia_ampliacion_eur,
        "revision": l.revision_puntos, "km_cert": _yn(l.km_certificados), "devolucion": l.devolucion_dias,
        "p_contado": l.precio_contado, "p_financiado": l.precio_financiado, "solo_fin": _yn(l.solo_financiando),
        "gastos": l.gastos_gestion_eur, "gastos_transf": _yn(l.gastos_incluyen_transferencia), "distancia": l.distancia_km,
        "duplicados": "; ".join(l.duplicados) or None, "avisos": avisos or None,
        "descripcion": (l.descripcion or "")[:1500], "fuente": f"{l.fuente_dato or ''} {l.raw_ref or ''}".strip(),
    }


def build_anuncios(wb: Workbook, listings: list[Listing], avisos: dict[str, str], names: Names, mkt=None) -> int:
    ws = wb.create_sheet("Anuncios", 0)
    n_last = max(len(listings) + 1, 2)
    for i, (key, hdr, kind, _t) in enumerate(COLUMNS, 1):
        c = ws.cell(1, i, hdr)
        c.font = Font(bold=True)
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if key in NOTES:
            c.comment = Comment(NOTES[key], "buscador")
    for r, l in enumerate(listings, 2):
        data = listing_row(l, avisos.get(l.uid, ""), (mkt or {}).get(l.uid))
        for i, (key, _h, kind, tpl) in enumerate(COLUMNS, 1):
            cell = ws.cell(r, i)
            if kind == "data":
                cell.value = data[key]
                cell.font = _font(BLUE)
                if key == "url" and l.url:
                    cell.hyperlink = l.url
            else:
                f, ext = _expand(tpl, r, n_last, names.names)
                ref = f"{COL[key]}{r}"
                cell.value = ArrayFormula(ref, f) if kind == "array" else f
                cell.font = _font(GREEN if ext else BLACK)
    widths = {"titulo": 40, "url": 30, "descripcion": 60, "avisos": 40, "duplicados": 30, "fuente": 30}
    for key, col in COL.items():
        ws.column_dimensions[col].width = widths.get(key, 13)
    return n_last


# =====================================================================================
# Top20
# =====================================================================================
TOP_COLS = [
    ("Título", "titulo"), ("Plataforma", "plataforma"), ("Marca", "marca"), ("Modelo", "modelo_key"),
    ("Motor", "motor"), ("Combustible", "combustible"), ("Año", "anio"), ("Km", "km"),
    ("Automático", "automatico"), ("Vendedor", "vendedor"), ("Precio contado (€)", "precio"),
    ("Solo financiando", "solo_fin"), ("Precio financiado (€)", "p_financiado"), ("Garantía efectiva (meses)", "gar_ef"),
    ("Gastos vendedor (€)", "gastos_calc"), ("ITP (€)", "itp"), ("Coste total (€)", "coste_total"), ("Mediana segmento (€)", "mediana"),
    ("Descuento", "descuento"), ("Etiqueta", "etiqueta"),
    ("P. Fiabilidad", "s_fiab"), ("P. Precio", "s_precio"), ("P. Km y edad", "s_kmedad"),
    ("P. Anuncio", "s_anuncio"), ("P. Coste uso", "s_uso"), ("P. Seguridad compra", "s_seguridad"),
    ("PUNTUACIÓN", "total"), ("Avisos", "avisos"),
]


def build_top(wb: Workbook, n_last: int, top_n: int) -> None:
    ws = wb.create_sheet("Top20", 1)
    hdrs = ["Puesto", "Fila en Anuncios", "Enlace"] + [h for h, _ in TOP_COLS] + \
           ["Aport. fiabilidad", "Aport. precio", "Aport. km/edad", "Aport. anuncio", "Aport. uso", "Aport. seguridad"]
    for i, h in enumerate(hdrs, 1):
        c = ws.cell(1, i, h)
        c.font = Font(bold=True)
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    rank_rng = f"Anuncios!${COL['rank']}$2:${COL['rank']}${n_last}"

    def rng(key):
        return f"Anuncios!${COL[key]}$2:${COL[key]}${n_last}"

    for k in range(1, top_n + 1):
        r = k + 1
        ws.cell(r, 1, k).font = _font(BLUE)
        ws.cell(r, 2, f'=IFERROR(MATCH(A{r},{rank_rng},0),"")').font = _font(GREEN)
        ws.cell(r, 3, f'=IF(B{r}="","",HYPERLINK(INDEX({rng("url")},B{r}),"Abrir anuncio"))').font = _font(GREEN)
        for j, (_h, key) in enumerate(TOP_COLS, 4):
            ws.cell(r, j, f'=IF($B{r}="","",INDEX({rng(key)},$B{r}))').font = _font(GREEN)
        base = 4 + len(TOP_COLS)
        sub = {"s_fiab": "W_FIAB", "s_precio": "W_PRECIO", "s_kmedad": "W_KMEDAD", "s_anuncio": "W_ANUNCIO", "s_uso": "W_USO",
               "s_seguridad": "W_SEGURIDAD"}
        for j, (key, w) in enumerate(sub.items()):
            colsrc = get_column_letter(4 + [x for _, x in TOP_COLS].index(key))
            ws.cell(r, base + j, f'=IF($B{r}="","",{colsrc}{r}*{w})').font = _font(GREEN)
    note = ws.cell(top_n + 3, 1, "Todo se calcula con fórmulas a partir de 'Anuncios' (verde) y 'Supuestos'. "
                                 "Aportación = puntuación del criterio × peso. Suma de aportaciones = PUNTUACIÓN.")
    note.font = Font(italic=True)
    for i in range(1, len(hdrs) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.column_dimensions["D"].width = 40


def write_workbook(path: Path, listings: list[Listing], avisos: dict[str, str], p: Params, mkt=None) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    names = Names(wb)
    build_supuestos(wb, p, names)
    n_last = build_anuncios(wb, listings, avisos, names, mkt)
    build_top(wb, n_last, p.cfg["salida"]["top_n"])
    for ws in wb.worksheets:
        ws.freeze_panes = None
    wb.save(path)
