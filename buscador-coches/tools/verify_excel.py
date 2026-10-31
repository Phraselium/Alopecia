"""Verifica ranking.xlsx: lo recalcula con LibreOffice y comprueba que no hay errores de fórmula, que no hay
paneles inmovilizados y que puntuaciones, costes y puestos coinciden con la réplica en Python."""
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE))
from src.excel import COL  # noqa: E402
from src.pipeline import process  # noqa: E402
from src.scoring import Params  # noqa: E402
from src.store import Store  # noqa: E402

L = lambda n: yaml.safe_load((BASE / n).read_text(encoding="utf-8"))  # noqa: E731
cfg = L("config.yaml")
p = Params(cfg, L("fiscal.yaml"), L("reliability.yaml"))
st = Store(BASE / cfg["salida"]["db"])
res = process(st.all(), p, set(), st.market())
with tempfile.TemporaryDirectory() as td:
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", td, str(BASE / cfg["salida"]["excel"])],
                   check=True, capture_output=True, timeout=300)
    wb = load_workbook(Path(td) / cfg["salida"]["excel"], data_only=True)
errors = [f"{sh.title}!{c.coordinate}={c.value}" for sh in wb.worksheets for row in sh.iter_rows() for c in row
          if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
frozen = [sh.title for sh in wb.worksheets if sh.freeze_panes]
ws = wb["Anuncios"]
hdr = {ws.cell(1, c).column_letter: c for c in range(1, ws.max_column + 1)}
mism = []
for r, f in enumerate(res.filas, 2):
    def v(k):
        return ws.cell(r, hdr[COL[k]]).value
    for k, exp in (("total", f["total"]), ("coste_total", f["coste"]["total"]), ("primer_anio", f["coste"]["primer_anio"])):
        if v(k) is None or abs(float(v(k)) - exp) > 1e-6:
            mism.append((f["listing"].uid, k, v(k), exp))
    if (v("rank") or None) != f.get("rank"):
        mism.append((f["listing"].uid, "rank", v("rank"), f.get("rank")))
print(f"filas={len(res.filas)} errores_formula={len(errors)} paneles_inmovilizados={frozen} discrepancias={len(mism)}")
for e in errors[:10] + mism[:10]:
    print("  ", e)
top = wb["Top20"]
print("Top 5 (Excel recalculado):", [top.cell(i, 4).value for i in range(2, 7)])
sys.exit(1 if errors or frozen or mism else 0)
