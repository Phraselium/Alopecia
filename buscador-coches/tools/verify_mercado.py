"""Verifica mercado.xlsx: lo recalcula con LibreOffice y comprueba 0 errores de fórmula, ningún panel inmovilizado y que
las columnas calculadas de Anuncios, las medianas/recuentos de 'Precios por modelo' y los recuentos de 'Oferta' coinciden con
Python. Uso: python tools/verify_mercado.py [ruta.xlsx]"""
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE))
from src.mercado.analysis import run_analysis  # noqa: E402
from src.mercado.excel import CHECKS, COL  # noqa: E402
from src.mercado.store import MercadoStore  # noqa: E402


def verify(xlsx: Path, res: dict) -> tuple[list, list, list]:
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", td, str(xlsx)],
                       check=True, capture_output=True, timeout=3600)
        wb = load_workbook(Path(td) / xlsx.name, data_only=True)
    errors = [f"{sh.title}!{c.coordinate}={c.value}" for sh in wb.worksheets for row in sh.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
    frozen = [sh.title for sh in wb.worksheets if sh.freeze_panes]
    ws = wb["Anuncios"]
    idx = {COL[k]: i for i, k in enumerate(COL)}
    mism = []
    for r, (a, row) in enumerate(zip(res["ads"], ws.iter_rows(min_row=2, values_only=True)), 2):
        if row[idx[COL["uid"]]] != a["uid"]:
            mism.append((r, "uid", row[0], a["uid"]))
            continue
        for k, g in CHECKS.items():
            exp, got = g(a), row[idx[COL[k]]]
            if isinstance(exp, str) or exp is None:
                if (got or None) != (exp or None):
                    mism.append((a["uid"], k, got, exp))
            elif got is None or got == "" or abs(float(got) - float(exp)) > 1e-6 * max(1, abs(float(exp))):
                mism.append((a["uid"], k, got, exp))
    pm = wb["Precios por modelo"]
    for row in pm.iter_rows(min_row=5, values_only=True):
        if row[0] == "modelo":
            if row[6] != row[2] or abs((row[7] or 0) - (row[3] or 0)) > 1e-6:
                mism.append(("precios", row[1], (row[6], row[7]), (row[2], row[3])))
    of = wb["Oferta"]
    for row in of.iter_rows(min_row=3, values_only=True):
        if isinstance(row[1], (int, float)) and isinstance(row[2], (int, float)) and row[1] != row[2]:
            mism.append(("oferta", row[0], row[2], row[1]))
    return errors, frozen, mism


if __name__ == "__main__":
    cfg = yaml.safe_load((BASE / "config_mercado.yaml").read_text(encoding="utf-8"))
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else BASE / cfg["salida"]["excel"]
    res = run_analysis(MercadoStore(BASE / cfg["salida"]["db"]).all(), cfg, BASE)
    errors, frozen, mism = verify(xlsx, res)
    print(f"filas={len(res['ads'])} errores_formula={len(errors)} paneles_inmovilizados={frozen} discrepancias={len(mism)}")
    for e in errors[:10] + mism[:15]:
        print("  ", e)
    sys.exit(1 if errors or frozen or mism else 0)
