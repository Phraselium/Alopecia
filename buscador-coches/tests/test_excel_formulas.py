"""Recalcula ranking.xlsx con LibreOffice y comprueba que las fórmulas dan lo mismo que src/scoring.py.

Usa anuncios SINTÉTICOS de prueba (no son datos reales y nunca se escriben en las salidas del proyecto).
"""
import random
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from src.excel import CHECKS, COL, liq_map, write_workbook
from src.models import Listing
from src.pipeline import process
from src.scoring import Params

pytestmark = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice no instalado")


def synthetic(n=80, seed=1):
    rnd = random.Random(seed)
    modelos = [("Toyota", "Yaris", "1.33 VVT-i"), ("Peugeot", "208", "1.2 PureTech"), ("Seat", "Ibiza", "1.4 MPI"),
               ("Renault", "Clio", "1.2 TCe"), ("Volkswagen", "Golf", "1.4 TSI"), ("Ford", "Focus", "1.6 TDCi diésel"),
               ("Kia", "Rio", "1.2 CVVT"), ("Skoda", "Octavia Combi", "1.6 TDI"), ("Opel", "Meriva", "1.4"),
               ("Dacia", "Logan", "1.4 MPI"), ("Peugeot", "207", "1.6 HDi"), ("Opel", "Corsa", "1.2 Easytronic"),
               ("Peugeot", "107", "1.0"), ("MINI", "MINI", "One 1.4"), ("Toyota", "Aygo", "1.0 MMT"), ("Opel", "Astra", "1.7 CDTi diésel"),
               ("Audi", "A4", "2.0 TFSI"), ("Toyota", "Avensis", "1.8 VVT-i"), ("Mazda", "6", "2.0"),
               ("Audi", "A3", "1.6 FSI"), ("Seat", "León", "1.6"), ("Alfa Romeo", "147", "1.6 TS"), ("Honda", "Civic", "1.8 i-VTEC")]
    colores = [None, "Blanco", "Gris / Plata", "COMMON.COLOR_BLUE", "Antracita", "Beige", "Rojo", "verde lima", "Plateado"]
    descs = ["ITV en vigor, libro de mantenimiento, color blanco", "único propietario, correa cambiada a los 120.000 km, facturas",
             "buen estado, llantas de aleación de 17\", faros xenón, lunas tintadas, ESP", "dos propietarios, tapacubos, neumáticos nuevos",
             "correa cambiada, libro de revisiones"]
    out = []
    for i in range(n):
        marca, modelo, motor = modelos[i % len(modelos)]
        anio = rnd.randint(2003, 2018)
        out.append(Listing(
            plataforma="TEST", id_externo=str(i),
            url=rnd.choice([f"https://example.invalid/{i}", f"https://example.invalid/x-gasolina-{anio - 1}-en-valencia-{i}-covo.aspx"]),
            titulo=f"{marca} {modelo} {motor}" + rnd.choice(["", " Sport", " FR"]), marca=marca, modelo=modelo, version=motor,
            combustible="diesel" if any(k in motor for k in ("TDCi", "TDI", "HDi", "CDTi")) else "gasolina",
            anio=anio, km=rnd.randint(40000, 219000), precio=float(rnd.randint(2500, 6000)),
            vendedor=rnd.choice(["particular", "profesional", None]), num_fotos=rnd.randint(1, 15),
            descripcion=rnd.choice(descs), color_raw=rnd.choice(colores),
            garantia_meses=rnd.choice([None, 12, 24]), garantia_ampliable=rnd.choice([None, True]),
            km_certificados=rnd.choice([None, True]), revision_puntos=rnd.choice([None, 150, 200]),
            devolucion_dias=rnd.choice([None, 7, 15]), gastos_gestion_eur=rnd.choice([None, 0.0, 350.0]),
            gastos_incluyen_transferencia=rnd.choice([None, True, False]),
            fecha_matriculacion=rnd.choice([None, f"{anio}", f"{anio}-{rnd.randint(1, 12):02d}-01", f"{rnd.randint(1, 12):02d}/{anio}",
                                           "2015-08-01", "2015", "2006"]),
            etiqueta_anuncio=rnd.choice([None, None, None, "C", "B", "ECO"])))
    return out


def test_formulas_match_python(tmp_path, cfg, fiscal, rel):
    # Las fórmulas se prueban con los filtros de vendedor/etiqueta desactivados para cubrir todos los casos
    cfg["filtros"].update({"solo_profesionales": False, "modelos_excluidos": [], "excluir_cajas_robotizadas": False,
                           "color": {"excluir_activo": False, "excluidos": []}})
    cfg["filtros"]["diesel"]["excluir_todos"] = False   # las fórmulas de diésel se siguen probando
    p = Params(cfg, fiscal, rel)
    ls = synthetic()
    market = [{"uid": f"m{i}", "plataforma": "TEST", "marca": l.marca, "modelo": l.modelo, "titulo": l.titulo, "anio": l.anio,
               "km": l.km, "precio": l.precio} for i, l in enumerate(ls * 2)]
    ventas = [{"marca": l.marca, "modelo": l.modelo, "titulo": l.titulo, "dias": 5 + 3 * i} for i, l in enumerate(ls[:30] * 3)]
    res = process(ls, p, set(), market, ventas=ventas)
    assert res.filas, "deberían quedar anuncios limpios"
    # el conjunto sintético debe ejercitar los casos nuevos
    assert any(f["penalizacion"] > 0 for f in res.filas), "falta carrocería penalizada"
    assert any(f["fiabilidad"]["diesel_pendiente"] for f in res.filas), "falta diésel sin fuente"
    assert any(f["estado"] == "negociar" for f in res.filas), "falta caso 'negociar'"
    assert any(f["listing"].etiqueta_dudosa for f in res.filas), "falta etiqueta dudosa"
    assert {f["listing"].etiqueta_fuente for f in res.filas} >= {"regla", "anuncio"}
    assert any(f["listing"].caja_robotizada for f in res.filas), "falta caja robotizada"
    assert any(f["fiabilidad"].get("gemelo") for f in res.filas), "falta gemelo de plataforma"
    assert any(f["listing"].etiqueta_dudosa_anio for f in res.filas), "falta dudosa por año"
    assert any(f["precio"]["tope"] < 100 for f in res.filas), "falta tope de precio por km"
    assert any(f["listing"].estetica_ref for f in res.filas), "falta modelo de estetica.yaml"
    assert len({f["color"]["score"] for f in res.filas}) >= 3, "faltan colores de varios grupos"
    assert any(f["listing"].at_llantas == "16+" for f in res.filas) and any(f["listing"].at_llantas == "chapa" for f in res.filas)
    assert any(f["listing"].ncap_estrellas for f in res.filas), "falta Euro NCAP"
    assert any(f["listing"].esp == "sí" for f in res.filas), "falta ESP"
    assert any(f["penal_presupuesto"] > 0 for f in res.filas), "falta penalización de presupuesto"
    assert any(f["coste"]["real_estimado"] for f in res.filas), "falta coste real de particular"
    assert any(f["reventa"]["n_dias"] >= 3 for f in res.filas), "faltan días de venta"
    assert any(f["seguridad"]["penal_particular"] for f in res.filas), "falta particular sin garantía"
    src = tmp_path / "r.xlsx"
    write_workbook(src, [f["listing"] for f in res.filas], res.avisos, p, {f["listing"].uid: f["mercado"] for f in res.filas}, liq_map(res.filas))
    outdir = tmp_path / "calc"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(outdir / "r.xlsx", data_only=True)
    ws = wb["Anuncios"]
    hdr = {ws.cell(1, c).column_letter: c for c in range(1, ws.max_column + 1)}
    # ningún error de fórmula en ninguna hoja
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith(("#", "Err:"))), f"{sh.title}!{c.coordinate}={c.value}"
    for r, f in enumerate(res.filas, 2):
        def v(key):
            return ws.cell(r, hdr[COL[key]]).value
        assert v("uid") == f["listing"].uid
        assert v("mediana") == pytest.approx(f["mediana"]), f["listing"].uid
        for k, getter in CHECKS.items():
            exp = getter(f)
            if exp is None or isinstance(exp, str):
                assert (v(k) or None) == (exp or None), (k, f["listing"].uid)
            else:
                assert v(k) == pytest.approx(exp), (k, f["listing"].uid)
        assert v("gastos_calc") == pytest.approx(f["coste"]["gastos"]), f["listing"].uid
        assert v("tasa") == pytest.approx(f["coste"]["tasa"]), f["listing"].uid
        assert v("itp") == pytest.approx(f["coste"]["itp"]), f["listing"].uid
        assert (v("rank") or None) == f.get("rank")
        l = f["listing"]
        assert (v("etq_dudosa") == "sí") == bool(l.etiqueta_dudosa), l.uid
    top = wb["Top20"]
    best = min((f for f in res.filas if f.get("rank")), key=lambda f: f["rank"])
    assert top["D2"].value == best["listing"].titulo
    for sh in wb.worksheets:
        assert sh.freeze_panes is None
